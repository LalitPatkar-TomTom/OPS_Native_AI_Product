"""Business impact report (Word) for OPS Native AI & OPS Hub, built only from the data on Azure.

Same structure as the 19 Oct business value report, but every number is computed from the Azure share:
    /data/config/native-ai/trigger_audit.db   analytics master (incl. the laptop history moved at go-live)
    /data/config/usage_log.jsonl, opshub_audit.db   OPS Hub questions and users
    /data/skills/Personal Skills/*_skill.md   profiles (scope)
It ends with a data-provenance appendix (tables, row counts, date ranges, migration marker) so the
reader can see where every number comes from.

    python impact_report.py --to lalit.patkar@tomtom.com        # build and email (Azure job)
    python impact_report.py --no-send --out x.docx --db ... --hub ... --skills ...   # build locally
Operations mail only: not recorded in the analytics DB.
"""
import argparse
import json
import re
import sqlite3
import sys
import tempfile
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from weekly_analytics import (CHAT_NOTE, REPORT_NOTE, chat_value_section, consolidated_value, report_value_section, IST, METRIC, MINUTES, PILOT_START, UC_COLOR, UC_NAME, INK2, GRID, NAVY,  # noqa: E402
                              ORANGE, AQUA, BLUE, chart_style, fill_breakdowns, ist_day, load, period_stats, table)
import matplotlib.pyplot as plt  # noqa: E402  (Agg backend set by weekly_analytics)
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

GREY = RGBColor(0x52, 0x52, 0x4E)
LIGHT = RGBColor(0x8A, 0x8A, 0x84)


def profiles_scope(skills: Path) -> dict:
    """Count profiles and distinct Databricks process types and planning IDs in *_skill.md files."""
    files = sorted(p for p in skills.glob("*_skill.md") if "@" in p.name)
    pts, pids, jira = set(), set(), set()
    for p in files:
        t = p.read_text(encoding="utf-8", errors="replace")
        for key, bag in (("Databricks Process Types", pts), ("Databricks Planning IDs", pids)):
            m = re.search(rf"\*\*{key}:\*\*\s*(.+)", t)
            if m:
                bag.update(v.strip().lower() for v in m.group(1).split(",")
                           if v.strip() and v.strip().lower() not in ("none", "-", "n/a"))
        for m in re.finditer(r"\(([A-Z][A-Z0-9]{1,9})\)", t[:4000]):
            jira.add(m.group(1))
    return {"profiles": len(files), "process_types": len(pts), "planning_ids": len(pids)}


def provenance(db: Path) -> list:
    """(table, rows, first day, last day) for each analytics table, read from a temp copy of the DB."""
    tmp = Path(tempfile.mkdtemp()) / "p.db"
    tmp.write_bytes(db.read_bytes())
    con = sqlite3.connect(f"file:{tmp}?mode=ro", uri=True)
    cols = {"trigger_events": "ts", "email_deliveries": "ts", "trigger_runs": "started_at",
            "uc4_alert_state": "alert_date", "history_log_daily": "day_ist"}
    out = []
    for t, c in cols.items():
        if not con.execute("SELECT 1 FROM sqlite_master WHERE name=?", (t,)).fetchone():
            continue
        n, lo, hi = con.execute(f'SELECT COUNT(*), MIN("{c}"), MAX("{c}") FROM "{t}"').fetchone()
        out.append((t, n, (lo or "")[:10], (hi or "")[:10]))
    con.close()
    return out


def main() -> None:
    """Build the business impact Word report from the Azure data and email it to --to via Graph.

    --no-send only builds it (to --out or a temp folder).
    """
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default="/data/config/native-ai/trigger_audit.db")
    ap.add_argument("--hub", default="/data/config")
    ap.add_argument("--skills", default="/data/skills/Personal Skills")
    ap.add_argument("--to", default="lalit.patkar@tomtom.com")
    ap.add_argument("--no-send", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()

    db = Path(a.db)
    d = load(db, Path(a.hub))
    now = datetime.now(IST)
    end = now.date()
    T = period_stats(d, PILOT_START, end)
    fill_breakdowns(T)
    days = [PILOT_START + timedelta(n) for n in range((end - PILOT_START).days + 1)]
    workdays = [x for x in days if x.weekday() < 5]
    scope = profiles_scope(Path(a.skills))
    users = scope["profiles"] or T["profiles"] or 1
    hub_users = d["hub_users"]
    marker = db.parent / "MIGRATED_FROM_LAPTOP.json"
    mig = json.loads(marker.read_text(encoding="utf-8")) if marker.exists() else {}
    prov = provenance(db)

    total_h = T["hours_total"]
    per_user_day = total_h / users / max(len(workdays), 1)
    annual_h = total_h / max(len(workdays), 1) * 220
    fte = annual_h / 1650
    alerts_by_metric = T["alerts_by_metric"]
    uc4_ev = [e for e in d["events"] if e["uc_id"] == "UC4"]
    alert_level = sum(1 for e in uc4_ev if e["status"] in ("sent", "suppressed"))
    by_day = Counter()
    for h in d.get("history", []):
        if h["uc_id"] != "RUN":
            by_day[(h["day_ist"], h["uc_id"])] += h["count"]
    for e in T["_sent"]:
        by_day[(ist_day(e["ts"]).isoformat(), e["uc_id"])] += 1
    runs_by_day = Counter()
    for h in d.get("history", []):
        if h["uc_id"] == "RUN":
            runs_by_day[h["day_ist"]] += h["count"]
    for r in d["runs"]:
        if r["trigger_source"] != "manual" and r["slot_date_ist"]:
            runs_by_day[r["slot_date_ist"]] += 1

    tmp = Path(tempfile.mkdtemp())
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    labels = [x.strftime("%d %b") for x in days]
    # fig 1 emails per day
    fig, ax = plt.subplots(figsize=(7.4, 3.0))
    bottom = [0] * len(days)
    for uc in ("UC1", "UC4", "UC2", "UC3"):
        vals = [by_day[(x.isoformat(), uc)] for x in days]
        ax.bar(labels, vals, bottom=bottom, color=UC_COLOR[uc], width=0.68, label=UC_NAME[uc], edgecolor="white", linewidth=1)
        bottom = [b + v for b, v in zip(bottom, vals)]
    for i, t in enumerate(bottom):
        if t:
            ax.text(i, t + 0.8, str(t), ha="center", va="bottom", fontsize=6.5, color=INK2)
    chart_style(ax, "Emails delivered")
    ax.tick_params(axis="x", rotation=70)
    ax.legend(frameon=False, fontsize=7.5, ncol=4, loc="upper left", bbox_to_anchor=(0, 1.18))
    c1 = tmp / "c1.png"; fig.savefig(c1, dpi=200, bbox_inches="tight", facecolor="white"); plt.close(fig)
    # fig 2 alerts by metric
    mk = sorted(alerts_by_metric.items(), key=lambda kv: -kv[1])
    fig, ax = plt.subplots(figsize=(5.2, 2.4))
    ax.bar([METRIC.get(k, k) for k, _ in mk], [v for _, v in mk], color=ORANGE, width=0.55)
    for i, (_, v) in enumerate(mk):
        ax.text(i, v + 0.5, str(v), ha="center", fontsize=8, color=INK2)
    chart_style(ax, "Breaches alerted")
    c2 = tmp / "c2.png"; fig.savefig(c2, dpi=200, bbox_inches="tight", facecolor="white"); plt.close(fig)
    # fig 3 value by activity
    acts = [(f"{UC_NAME[uc]} ({uc})", T["hours"][uc]) for uc in ("UC1", "UC4", "UC2", "UC3")] + \
           [("OPS Hub questions answered", T["hours"]["chat"]), ("OPS Hub Status Reports", T["hours"]["reports"])]
    acts.sort(key=lambda x: x[1])
    fig, ax = plt.subplots(figsize=(7.2, 2.5))
    ax.barh([x[0] for x in acts], [x[1] for x in acts], color=BLUE, height=0.55)
    for i, (_, h) in enumerate(acts):
        ax.text(h + 0.6, i, f"{h:.0f} h", va="center", fontsize=8, color=INK2)
    chart_style(ax)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8); ax.yaxis.grid(False)
    c3 = tmp / "c3.png"; fig.savefig(c3, dpi=200, bbox_inches="tight", facecolor="white"); plt.close(fig)
    # fig 4 runs per day
    fig, ax = plt.subplots(figsize=(7.4, 2.4))
    ax.bar(labels, [runs_by_day.get(x.isoformat(), 0) for x in days], color=AQUA, width=0.6)
    ax.axhline(4, color=INK2, linewidth=1, linestyle=(0, (4, 3)))
    ax.text(len(days) - 0.5, 4.1, "4 planned slots / day", ha="right", fontsize=7.5, color=INK2)
    chart_style(ax, "Runs completed"); ax.set_ylim(0, 4.8); ax.tick_params(axis="x", rotation=70)
    c4 = tmp / "c4.png"; fig.savefig(c4, dpi=200, bbox_inches="tight", facecolor="white"); plt.close(fig)

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = "Segoe UI"; st.font.size = Pt(10)
    for lvl, size in ((1, 15), (2, 12)):
        h = doc.styles[f"Heading {lvl}"]; h.font.name = "Segoe UI"; h.font.size = Pt(size)
        h.font.color.rgb = NAVY; h.font.bold = True

    def para(text, size=10, bold=False, color=None, italic=False, after=4):
        """Add a paragraph with one formatted run and return it."""
        p = doc.add_paragraph(); r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color:
            r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)
        return p

    def figure(path, caption, width=16.5):
        """Add a centred picture with an italic caption below it."""
        doc.add_picture(str(path), width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para(caption, size=8, italic=True, color=LIGHT, after=8).alignment = WD_ALIGN_PARAGRAPH.CENTER

    rng = f"{PILOT_START:%d %b} – {end:%d %b %Y}"
    para("", after=50)
    para("OPS Native AI & OPS Hub", size=26, bold=True, color=NAVY, after=4)
    para(f"Business Impact Report — {rng}", size=15, color=RGBColor(0x2A, 0x78, 0xD6), after=20)
    para(f"Built on Azure from the OPS Native AI analytics database · {now:%d %b %Y %H:%M} IST", size=10, color=GREY)
    para("TEST EDITION — check of the Azure data after the switch-over; numbers are complete for the pilot to date.",
         size=9, italic=True, color=LIGHT, after=30)

    doc.add_heading("1. Summary", 1)
    table(doc, ["Measure", "Value"], [
        ("Profile owners live", users),
        ("Emails delivered", f"{T['emails']:,}"),
        ("Email delivery success", f"{T['emails']}/{T['emails_tried']}"),
        ("Triggers run (all use cases)", f"{T['triggers']:,}"),
        ("Metric checks run", f"{T['checks']:,}"),
        ("Metric breaches alerted", T["alerts"]),
        ("Repeat alerts held back", T["held_back"]),
        ("Jira tickets flagged for SLA", T["tickets_flagged"]),
        ("OPS Hub questions answered", T["chats"]),
        ("OPS Hub Status Reports generated", T["reports"]),
        ("OPS Hub registered users", len(hub_users)),
        ("Scheduled runs done / planned", f"{T['runs']}/{T['slots_planned']}"),
        ("Estimated manual work avoided", f"{total_h:.0f} h  (≈{per_user_day * 60:.0f} min per user per working day)"),
        ("Annualised at today's users", f"≈{annual_h:,.0f} h a year ≈ {fte:.1f} FTE  (half the assumed effort: {fte / 2:.1f} FTE)"),
    ], widths=(7.0, 9.5), num_cols=())

    doc.add_heading("2. Scope", 1)
    table(doc, ["Measure", "Value"], [
        ("Profiles on the Azure share", scope["profiles"]),
        ("Process types monitored", scope["process_types"]),
        ("Planning IDs monitored", scope["planning_ids"]),
        ("People who received emails", T["people"]),
        ("OPS Hub users who asked questions", T["chat_users"]),
    ], widths=(7.0, 9.5))

    doc.add_heading("3. Delivery volume", 1)
    figure(c1, f"Figure 1 — Emails delivered per day by use case, {rng} (total {T['emails']:,}).")
    table(doc, ["Use case", "Emails", "Share"], [
        (UC_NAME[uc], T["by_uc"].get(uc, 0), f"{T['by_uc'].get(uc, 0) / max(T['emails'], 1):.0%}") for uc in ("UC1", "UC4", "UC2", "UC3")
    ] + [("Total", T["emails"], "100%")], widths=(7, 4, 3), num_cols=(1, 2))

    doc.add_heading("4. Early warning on quality and efficiency (UC4)", 1)
    para(f"{T['checks']:,} metric checks against each owner's own thresholds; {T['alerts']} distinct breaches "
         f"(one per metric per user per day) were alerted. {T['held_back']} repeat alerts were held back "
         f"({T['held_back'] / max(alert_level, 1):.0%} of alert-level results), so each owner gets one actionable message.")
    figure(c2, "Figure 2 — Breaches alerted by metric.", width=11)

    doc.add_heading("5. Jira SLA governance (UC2)", 1)
    para(f"{T['sla_checks']} SLA checks flagged {T['tickets_flagged']} distinct Jira tickets past their SLA age, "
         f"with {T['by_uc'].get('UC2', 0)} follow-up emails listing them.")

    doc.add_heading("6. Productivity value (estimate)", 1)
    figure(c3, f"Figure 3 — Estimated manual work avoided by activity (total {total_h:.0f} h).")
    chat_value_section(doc, T["chat_value"], "pilot")
    para(CHAT_NOTE, size=8.5, italic=True, color=GREY)
    report_value_section(doc, T["report_value"], "pilot")
    para(REPORT_NOTE, size=8.5, italic=True, color=GREY)
    consolidated_value(doc, T, "pilot")
    para(f"Annualised = pilot rate over {len(workdays)} working days × 220 working days; 1 FTE = 1,650 h. "
         "Effort per item is an assumption to be validated with users (OSI-023).", size=8.5, italic=True, color=GREY)

    doc.add_heading("7. Reliability", 1)
    figure(c4, f"Figure 4 — Runs completed per day against 4 planned slots ({T['runs']} of {T['slots_planned']}).")
    para(f"Run errors recorded: {T['run_errors']}. Runs more than 30 min late: {T['late_runs']}. "
         "Since 20 Oct the engine runs as an Azure job (opsai@tomtom.com), no longer on a laptop.")

    doc.add_heading("8. Per profile owner", 1)
    owners = sorted(T["per_owner_uc"].items(), key=lambda kv: -sum(kv[1].values()))
    table(doc, ["Profile owner", "Briefings", "Alerts", "SLA", "Weekly", "Total"],
          [(o, c.get("UC1", 0), c.get("UC4", 0), c.get("UC2", 0), c.get("UC3", 0), sum(c.values())) for o, c in owners],
          widths=(7.0, 1.9, 1.7, 1.6, 1.7, 1.6), num_cols=(1, 2, 3, 4, 5))
    para("Per-owner counts cover emails recorded in the analytics database (from 13 Oct); 1–13 Oct emails are in the "
         "totals from the run-log history.", size=8.5, italic=True, color=GREY)

    doc.add_heading("Appendix — data provenance", 1)
    para(f"Analytics database: {db} (Azure Files share opshub-config).")
    if mig:
        para(f"Laptop history moved to Azure at {mig.get('migrated_at_utc', '?')} UTC (checksum {str(mig.get('sha256', ''))[:12]}…); "
             f"from then on this database is the only analytics master.")
    table(doc, ["Table", "Rows", "From", "To"], [(t, f"{n:,}", lo, hi) for t, n, lo, hi in prov],
          widths=(6.0, 2.5, 3.5, 3.5), num_cols=(1,))
    table(doc, ["Other source on Azure", "Used for"], [
        ("opshub-config/usage_log.jsonl", f"OPS Hub questions ({len(d['chats'])} lines)"),
        ("opshub-config/opshub_audit.db", f"OPS Hub registered users ({len(hub_users)})"),
        ("opshub-skills/Personal Skills", f"Profiles and monitored scope ({scope['profiles']} files)"),
    ], widths=(7.0, 9.5))
    for t in ("Test sends are excluded. Days are IST.",
              "Breach = one metric over the owner's alert threshold, counted once per user per day.",
              "Hours saved are estimates from stated assumptions, not time tracking."):
        doc.add_paragraph(t, style="List Bullet")

    out = Path(a.out) if a.out else Path(tempfile.mkdtemp()) / f"OPS_NativeAI_OPSHub_Business_Impact_{end:%Y-%m-%d}_Azure.docx"
    doc.save(out)
    print(f"impact report built: {out.name} — {T['emails']} emails, {T['alerts']} breaches, {total_h:.0f} h, provenance {prov}")
    if a.no_send:
        return
    sys.path.insert(0, "/app/MultiAgent_Briefing")
    from delivery import _send_graph
    html = (f'<div style="font-family:Segoe UI,Arial,sans-serif;font-size:13px"><p>Test edition of the business impact '
            f"report, built on Azure from the OPS Native AI analytics database ({rng}).</p><ul>"
            f"<li>{T['emails']:,} emails delivered ({T['emails']}/{T['emails_tried']} accepted)</li>"
            f"<li>{T['triggers']:,} triggers, {T['alerts']} breaches alerted, {T['tickets_flagged']} SLA tickets flagged</li>"
            f"<li>{T['chats']} OPS Hub questions; {len(hub_users)} registered users</li>"
            f"<li>≈{total_h:.0f} h of manual work avoided (estimate), ≈{fte:.1f} FTE annualised</li></ul>"
            "<p>The appendix lists every table, row count and date range used.</p></div>")
    ok = _send_graph(a.to, f"[TEST] OPS Native AI & OPS Hub — Business Impact Report {rng} (from Azure data)", html, None, str(out))
    print(f"impact report {'sent to ' + a.to if ok else 'NOT sent'}")


if __name__ == "__main__":
    main()

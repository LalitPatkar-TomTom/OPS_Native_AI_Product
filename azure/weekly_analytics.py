"""Weekly OPS Native AI & OPS Hub analytics (Word) — every Monday after the 09:00 run, from opsai@tomtom.com,
To the four report readers, CC Peter Clarys, Seema Nayyar and Saurav Nanda.

Covers the previous Monday–Sunday, compared with the week before and the total since the pilot start.
Read-only on all data. Sources (Azure share, mounted at /data/config):
    native-ai/trigger_audit.db   OPS Native AI analytics DB (the master after the switch-over)
    usage_log.jsonl              OPS Hub questions asked (one line per question)
    opshub_audit.db              OPS Hub registered users

    python weekly_analytics.py                       # called by entrypoint.sh on Mondays, 09:00 slot
    python weekly_analytics.py --to me@tomtom.com    # test: send only to this address
    python weekly_analytics.py --no-send --out x.docx --db <path> --hub <dir>   # build locally
The report email is an operations mail: it is not recorded in the analytics DB.
"""
import argparse
import json
import os
import shutil
import sqlite3
import sys
import tempfile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

RECIPIENTS = ["Ketki.Avachat@tomtom.com", "Vrushali.Kashikar@tomtom.com",
              "pratish.deshpande@tomtom.com", "lalit.patkar@tomtom.com"]          # To
CC = ["Peter.Clarys@tomtom.com", "Seema.Nayyar2@tomtom.com", "Saurav.Nanda@tomtom.com"]
IST = timezone(timedelta(hours=5, minutes=30))
PILOT_START = date(2026, 10, 1)

UC_NAME = {"UC1": "Daily Briefing", "UC2": "SLA Follow-up", "UC3": "Weekly Report", "UC4": "Metric Alert"}
BLUE, ORANGE, AQUA, YELLOW = "#2a78d6", "#eb6834", "#1baf7a", "#eda100"
UC_COLOR = {"UC1": BLUE, "UC4": ORANGE, "UC2": AQUA, "UC3": YELLOW}
INK2, GRID = "#52524e", "#e6e6e1"
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
METRIC = {"coq": "CoQ", "copq": "CoPQ", "efficiency": "Efficiency", "yield": "Yield", "fta": "FTA"}
# minutes of manual work avoided per item — the same assumptions as the 19 Oct business value report
MINUTES = {"UC1": 20, "UC4": 15, "UC2": 10, "UC3": 90, "chat": 10}


def ist_day(ts: str) -> date:
    """ISO timestamp (UTC or with offset) -> IST calendar day."""
    d = datetime.fromisoformat(ts.replace("Z", "+00:00"))
    if d.tzinfo is None:
        d = d.replace(tzinfo=timezone.utc)
    return d.astimezone(IST).date()


def load(db_path: Path, hub_dir: Path) -> dict:
    tmp = Path(tempfile.mkdtemp())
    local = tmp / "a.db"                         # never read SQLite over SMB in place
    shutil.copy2(db_path, local)
    con = sqlite3.connect(f"file:{local}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    q = lambda sql: [dict(r) for r in con.execute(sql)]
    data = {
        "emails": q("SELECT ts, uc_id, profile_owner, to_email, sent_ok, has_attachment FROM email_deliveries WHERE is_test = 0"),
        "events": q("SELECT ts, uc_id, user_email, status, metric_key, sla_breaches, sla_warnings, "
                    "sla_breach_tickets, sla_warning_tickets FROM trigger_events WHERE is_test = 0"),
        "runs": q("SELECT slot_date_ist, slot_ist, delay_min, errors, trigger_source FROM trigger_runs WHERE is_test = 0"),
        "alerts": q("SELECT user_email, metric_key, alert_date FROM uc4_alert_state"),
        # emails and runs before the DB tables started (1–13 Oct), moved from the laptop log at the switch-over
        "history": q("SELECT day_ist, uc_id, count FROM history_log_daily")
        if con.execute("SELECT 1 FROM sqlite_master WHERE name = 'history_log_daily'").fetchone() else [],
    }
    con.close()
    chats = []
    if (hub_dir / "usage_log.jsonl").exists():
        for line in (hub_dir / "usage_log.jsonl").read_text(encoding="utf-8", errors="replace").splitlines():
            try:
                chats.append(json.loads(line))
            except ValueError:
                pass
    data["chats"] = chats
    users = []
    if (hub_dir / "opshub_audit.db").exists():
        shutil.copy2(hub_dir / "opshub_audit.db", tmp / "h.db")
        h = sqlite3.connect(f"file:{tmp / 'h.db'}?mode=ro", uri=True)
        users = [dict(zip(("email", "added_at", "last_active"), r))
                 for r in h.execute("SELECT email, added_at, last_active FROM opshub_users")]
        h.close()
    data["hub_users"] = users
    return data


def tickets(e: dict) -> set:
    out = set()
    for col in ("sla_breach_tickets", "sla_warning_tickets"):
        try:
            for t in json.loads(e.get(col) or "[]"):
                out.add(t if isinstance(t, str) else (t.get("key") if isinstance(t, dict) else str(t)))
        except ValueError:
            pass
    return out


def period_stats(d: dict, start: date, end: date) -> dict:
    """Numbers for start..end (inclusive, IST days)."""
    inside = lambda ts: start <= ist_day(ts) <= end
    em = [e for e in d["emails"] if inside(e["ts"])]
    ev = [e for e in d["events"] if inside(e["ts"])]
    runs = [r for r in d["runs"] if r["slot_date_ist"] and start.isoformat() <= r["slot_date_ist"] <= end.isoformat()
            and r["trigger_source"] != "manual"]
    alerts = [a for a in d["alerts"] if start.isoformat() <= a["alert_date"] <= end.isoformat()
              and "test.user" not in (a["user_email"] or "")]
    chats = [c for c in d["chats"] if c.get("ts") and inside(c["ts"])]
    sent = [e for e in em if e["sent_ok"]]
    by_uc = Counter(e["uc_id"] for e in sent)
    hist = [h for h in d.get("history", []) if start.isoformat() <= h["day_ist"] <= end.isoformat()]
    hist_emails = Counter()
    for h in hist:
        if h["uc_id"] != "RUN":
            hist_emails[h["uc_id"]] += h["count"]
    by_uc.update(hist_emails)
    hist_runs = sum(h["count"] for h in hist if h["uc_id"] == "RUN")
    uc2 = [e for e in ev if e["uc_id"] == "UC2"]
    flagged = set().union(*[tickets(e) for e in uc2]) if uc2 else set()
    days = (end - start).days + 1
    slots = {(r["slot_date_ist"], r["slot_ist"]) for r in runs}
    hours = {uc: by_uc.get(uc, 0) * MINUTES[uc] / 60 for uc in ("UC1", "UC2", "UC3", "UC4")}
    hours["chat"] = len(chats) * MINUTES["chat"] / 60
    return {
        "emails": len(sent) + sum(hist_emails.values()), "emails_tried": len(em) + sum(hist_emails.values()),
        "by_uc": by_uc,
        "people": len({e["to_email"].lower() for e in sent if e["to_email"]}),
        "profiles": len({e["user_email"].lower() for e in ev if e["user_email"]}),
        "triggers": len(ev), "triggers_by_uc": Counter(e["uc_id"] for e in ev),
        "checks": sum(1 for e in ev if e["uc_id"] == "UC4"),
        "alerts": len(alerts), "alerts_by_metric": Counter(a["metric_key"] for a in alerts),
        "held_back": sum(1 for e in ev if e["uc_id"] == "UC4" and e["status"] == "suppressed"),
        "sla_checks": len(uc2), "tickets_flagged": len(flagged),
        "runs": len(slots) + hist_runs, "slots_planned": 4 * days,
        "late_runs": sum(1 for r in runs if (r["delay_min"] or 0) > 30),
        "run_errors": sum(r["errors"] or 0 for r in runs),
        "chats": len(chats), "chat_users": len({c.get("email", "").lower() for c in chats if c.get("email")}),
        "hours": hours, "hours_total": sum(hours.values()),
        "per_owner": Counter((e["profile_owner"] or e["to_email"] or "").lower() for e in sent),
        "per_owner_uc": defaultdict(Counter),
        "by_day_uc": defaultdict(Counter),
        "_sent": sent,
    }


def fill_breakdowns(s: dict) -> None:
    for e in s["_sent"]:
        s["per_owner_uc"][(e["profile_owner"] or e["to_email"] or "").lower()][e["uc_id"]] += 1
        s["by_day_uc"][ist_day(e["ts"])][e["uc_id"]] += 1


# ── document helpers ─────────────────────────────────────────────────────────
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_fill)
    tcPr.append(sh)


def table(doc, headers, rows, widths=None, num_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1B2A4A")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(v)); run.font.size = Pt(9)
            if i in num_cols:
                p.alignment = 2
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def chart_style(ax, ylabel=""):
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color(GRID)
    ax.tick_params(colors=INK2, labelsize=8, length=0)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    if ylabel:
        ax.set_ylabel(ylabel, color=INK2, fontsize=8.5)


def fmt_change(now, before):
    if before == 0:
        return "new" if now else "–"
    pct = (now - before) / before * 100
    return f"{pct:+.0f}%"


def build(d: dict, week_start: date, out: Path, source: str) -> tuple[Path, dict]:
    week_end = week_start + timedelta(days=6)
    prev_start = week_start - timedelta(days=7)
    W = period_stats(d, week_start, week_end); fill_breakdowns(W)
    P = period_stats(d, prev_start, week_start - timedelta(days=1))
    T = period_stats(d, PILOT_START, week_end)
    hub_users = d["hub_users"]
    new_hub = sum(1 for u in hub_users if u["added_at"] and week_start <= ist_day(u["added_at"]) <= week_end)
    active_hub = sum(1 for u in hub_users if u["last_active"] and week_start <= ist_day(u["last_active"]) <= week_end)
    tmp = Path(tempfile.mkdtemp())
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})

    # chart 1: emails per day by use case
    days = [week_start + timedelta(n) for n in range(7)]
    fig, ax = plt.subplots(figsize=(7.0, 2.8))
    labels = [x.strftime("%a %d %b") for x in days]
    bottom = [0] * 7
    for uc in ("UC1", "UC4", "UC2", "UC3"):
        vals = [W["by_day_uc"][x].get(uc, 0) for x in days]
        if any(vals):
            ax.bar(labels, vals, bottom=bottom, color=UC_COLOR[uc], width=0.62, label=UC_NAME[uc],
                   edgecolor="white", linewidth=1.2)
        bottom = [b + v for b, v in zip(bottom, vals)]
    for i, t in enumerate(bottom):
        if t:
            ax.text(i, t + 0.6, str(t), ha="center", va="bottom", fontsize=7.5, color=INK2)
    chart_style(ax, "Emails delivered")
    ax.legend(frameon=False, fontsize=7.5, ncol=4, loc="upper left", bbox_to_anchor=(0, 1.2))
    c1 = tmp / "c1.png"; fig.savefig(c1, dpi=200, bbox_inches="tight", facecolor="white"); plt.close(fig)

    # chart 2: alerts by metric
    c2 = None
    if W["alerts_by_metric"]:
        mk = sorted(W["alerts_by_metric"].items(), key=lambda kv: -kv[1])
        fig, ax = plt.subplots(figsize=(5.0, 2.3))
        ax.bar([METRIC.get(k, k) for k, _ in mk], [v for _, v in mk], color=ORANGE, width=0.55)
        for i, (_, v) in enumerate(mk):
            ax.text(i, v + 0.3, str(v), ha="center", fontsize=8, color=INK2)
        chart_style(ax, "Alerts raised")
        c2 = tmp / "c2.png"; fig.savefig(c2, dpi=200, bbox_inches="tight", facecolor="white"); plt.close(fig)

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = "Segoe UI"; st.font.size = Pt(10)
    for lvl, size in ((1, 15), (2, 12)):
        h = doc.styles[f"Heading {lvl}"]; h.font.name = "Segoe UI"; h.font.size = Pt(size)
        h.font.color.rgb = NAVY; h.font.bold = True

    rng = f"{week_start:%d %b} – {week_end:%d %b %Y}"
    doc.add_heading("OPS Native AI & OPS Hub — Weekly Analytics", 1)
    p = doc.add_paragraph(f"Week {week_start.isocalendar()[1]}: {rng}  ·  compared with the week before and the pilot total "
                          f"since {PILOT_START:%d %b %Y}")
    p.runs[0].font.color.rgb = RGBColor(0x52, 0x52, 0x4E)

    doc.add_heading("Key numbers", 2)
    rows = [
        ("Triggers run (all use cases)", W["triggers"], P["triggers"], fmt_change(W["triggers"], P["triggers"]), T["triggers"]),
        ("Profiles served", W["profiles"], P["profiles"], fmt_change(W["profiles"], P["profiles"]), T["profiles"]),
        ("Emails delivered", W["emails"], P["emails"], fmt_change(W["emails"], P["emails"]), T["emails"]),
        ("Delivery success", f"{W['emails']}/{W['emails_tried']}", f"{P['emails']}/{P['emails_tried']}", "", f"{T['emails']}/{T['emails_tried']}"),
        ("People who received emails", W["people"], P["people"], fmt_change(W["people"], P["people"]), T["people"]),
        ("Daily briefings (UC1)", W["by_uc"].get("UC1", 0), P["by_uc"].get("UC1", 0), fmt_change(W["by_uc"].get("UC1", 0), P["by_uc"].get("UC1", 0)), T["by_uc"].get("UC1", 0)),
        ("Metric alerts sent (UC4)", W["by_uc"].get("UC4", 0), P["by_uc"].get("UC4", 0), fmt_change(W["by_uc"].get("UC4", 0), P["by_uc"].get("UC4", 0)), T["by_uc"].get("UC4", 0)),
        ("SLA follow-ups (UC2)", W["by_uc"].get("UC2", 0), P["by_uc"].get("UC2", 0), fmt_change(W["by_uc"].get("UC2", 0), P["by_uc"].get("UC2", 0)), T["by_uc"].get("UC2", 0)),
        ("Weekly reports (UC3)", W["by_uc"].get("UC3", 0), P["by_uc"].get("UC3", 0), fmt_change(W["by_uc"].get("UC3", 0), P["by_uc"].get("UC3", 0)), T["by_uc"].get("UC3", 0)),
        ("Metric checks run", W["checks"], P["checks"], fmt_change(W["checks"], P["checks"]), T["checks"]),
        ("Repeat alerts held back", W["held_back"], P["held_back"], "", T["held_back"]),
        ("Jira tickets flagged for SLA", W["tickets_flagged"], P["tickets_flagged"], fmt_change(W["tickets_flagged"], P["tickets_flagged"]), T["tickets_flagged"]),
        ("Scheduled runs done / planned", f"{W['runs']}/{W['slots_planned']}", f"{P['runs']}/{P['slots_planned']}", "", f"{T['runs']}/{T['slots_planned']}"),
        ("OPS Hub questions answered", W["chats"], P["chats"], fmt_change(W["chats"], P["chats"]), T["chats"]),
        ("OPS Hub users asking questions", W["chat_users"], P["chat_users"], "", T["chat_users"]),
        ("Estimated hours of manual work avoided", f"{W['hours_total']:.0f} h", f"{P['hours_total']:.0f} h", fmt_change(W["hours_total"], P["hours_total"]), f"{T['hours_total']:.0f} h"),
    ]
    table(doc, ["Measure", "This week", "Week before", "Change", "Pilot total"], rows,
          widths=(6.4, 2.4, 2.4, 1.9, 2.4), num_cols=(1, 2, 3, 4))
    doc.add_paragraph(f"OPS Hub: {len(hub_users)} registered users, {new_hub} added and {active_hub} active this week.")

    doc.add_heading("Emails per day", 2)
    doc.add_picture(str(c1), width=Cm(16.5))
    if c2:
        doc.add_heading("Metric alerts by metric", 2)
        doc.add_picture(str(c2), width=Cm(11))

    doc.add_heading("Per profile owner (this week)", 2)
    owners = sorted(W["per_owner_uc"].items(), key=lambda kv: -sum(kv[1].values()))
    table(doc, ["Profile owner", "Briefings", "Alerts", "SLA", "Weekly", "Total"],
          [(o, c.get("UC1", 0), c.get("UC4", 0), c.get("UC2", 0), c.get("UC3", 0), sum(c.values())) for o, c in owners]
          or [("No emails this week", "", "", "", "", "")],
          widths=(7.0, 1.9, 1.7, 1.6, 1.7, 1.6), num_cols=(1, 2, 3, 4, 5))

    doc.add_heading("Estimated value this week", 2)
    vrows = [(f"{UC_NAME[uc]} ({uc})", W["by_uc"].get(uc, 0), MINUTES[uc], f"{W['hours'][uc]:.1f} h") for uc in ("UC1", "UC4", "UC2", "UC3")]
    vrows.append(("OPS Hub questions answered", W["chats"], MINUTES["chat"], f"{W['hours']['chat']:.1f} h"))
    vrows.append(("Total", "", "", f"{W['hours_total']:.1f} h"))
    table(doc, ["Activity", "Items", "Minutes saved per item (assumption)", "Hours"], vrows,
          widths=(6.5, 2.0, 4.6, 2.4), num_cols=(1, 2, 3))

    doc.add_heading("Reliability", 2)
    late = f", {W['late_runs']} more than 30 min late" if W["late_runs"] else ""
    doc.add_paragraph(f"{W['runs']} of {W['slots_planned']} scheduled runs done this week{late}; "
                      f"{W['run_errors']} run error(s); {W['emails']} of {W['emails_tried']} emails accepted by the mail service.")

    doc.add_heading("About these numbers", 2)
    for t in (f"Source: {source}. Test sends are excluded. Days are IST.",
              "Pilot totals include 1–13 Oct emails and runs counted from the laptop run log (before the analytics "
              "database recorded them); per-recipient detail and alert checks exist from 9–13 Oct onwards.",
              "Hours saved use the same per-item assumptions as the 19 Oct business value report; they are estimates "
              "to be validated with users (OSI-023).",
              "OPS Hub questions come from the OPS Hub usage log; Status Report volumes are not in it yet.",
              f"Generated {datetime.now(IST):%d %b %Y %H:%M} IST by the Azure job ops-native-ai-trigger."):
        doc.add_paragraph(t, style="List Bullet")

    doc.save(out)
    return out, {"W": W, "P": P, "T": T, "rng": rng}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default=os.getenv("OPS_DATA_DIR", "/data/config/native-ai") + "/trigger_audit.db")
    ap.add_argument("--hub", default="/data/config")
    ap.add_argument("--week-start", help="Monday of the week to report (default: last week)")
    ap.add_argument("--to", help="send only to this address (test)")
    ap.add_argument("--no-send", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()

    today = datetime.now(IST).date()
    ws = date.fromisoformat(a.week_start) if a.week_start else today - timedelta(days=today.weekday() + 7)
    marker = Path(a.db).parent / "MIGRATED_FROM_LAPTOP.json"
    source = ("OPS Native AI analytics database on Azure (includes the laptop history moved at the switch-over)"
              if marker.exists() else "OPS Native AI analytics database")
    out = Path(a.out) if a.out else Path(tempfile.mkdtemp()) / f"OPS_NativeAI_OPSHub_Weekly_Analytics_{ws:%Y-%m-%d}.docx"
    path, s = build(load(Path(a.db), Path(a.hub)), ws, out, source)
    W = s["W"]
    print(f"weekly analytics built: {path.name} ({W['emails']} emails, {W['alerts']} alerts, {W['chats']} OPS Hub questions)")
    if a.no_send:
        return

    sys.path.insert(0, "/app/MultiAgent_Briefing")
    from delivery import _send_graph
    to = a.to or ";".join(RECIPIENTS)
    html = (f'<div style="font-family:Segoe UI,Arial,sans-serif;font-size:13px">'
            f"<p>Hello,</p><p>Attached is the weekly OPS Native AI &amp; OPS Hub analytics for <b>{s['rng']}</b>.</p>"
            f"<ul><li>{W['triggers']} triggers run for {W['profiles']} profiles</li>"
            f"<li>{W['emails']} emails delivered to {W['people']} people ({W['by_uc'].get('UC1', 0)} briefings, "
            f"{W['by_uc'].get('UC4', 0)} metric alerts, {W['by_uc'].get('UC2', 0)} SLA follow-ups, "
            f"{W['by_uc'].get('UC3', 0)} weekly reports)</li>"
            f"<li>{W['alerts']} metric breaches alerted, {W['held_back']} repeat alerts held back</li>"
            f"<li>{W['chats']} OPS Hub questions answered for {W['chat_users']} users</li>"
            f"<li>About {W['hours_total']:.0f} hours of manual work avoided (estimate)</li>"
            f"<li>{W['runs']} of {W['slots_planned']} scheduled runs done</li></ul>"
            f'<p style="color:#888;font-size:11px">Sent automatically every Monday by OPS Native AI from opsai@tomtom.com.</p></div>')
    cc = None if a.to else CC                  # a test (--to) goes to that address only
    ok = _send_graph(to, f"OPS Native AI & OPS Hub — Weekly Analytics {s['rng']}", html, cc, str(path))
    print(f"weekly analytics {'sent to ' + to + (' cc ' + ', '.join(cc) if cc else '') if ok else 'NOT sent'}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:        # never fail the run because of the report
        print(f"weekly analytics failed: {type(exc).__name__}: {exc}")

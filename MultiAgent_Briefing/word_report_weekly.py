"""
Word (.docx) report builder for UC3 Weekly Operational Report.
Mirrors the HTML email content — used as the email attachment.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any


def build(
    jira: dict[str, Any],
    analytics: dict[str, Any],
    confluence: dict[str, Any] | None = None,
    skill: dict[str, Any] | None = None,
) -> bytes:
    """
    Build a Word report from the same data as the weekly HTML email.
    Returns raw .docx bytes.
    Raises RuntimeError if python-docx is not installed.
    """
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")

    _s           = skill or {}
    project_name = _s.get("project_name") or "Project"
    jira_project = _s.get("jira_project") or ""
    user_name    = _s.get("user_first_name") or "you"
    pm           = set(_s.get("primary_metrics") or ["fta", "efficiency", "coq"])
    proj_label   = f"{project_name} ({jira_project})" if jira_project else project_name

    confluence   = confluence or {}
    week         = analytics.get("week_label", datetime.now().strftime("%Y-%m-%d"))
    now          = datetime.now(timezone.utc)
    date_str     = now.strftime("%A, %d %B %Y  ·  %H:%M UTC")
    m4w          = list(reversed(analytics.get("metrics_4w", [])))
    breakdown    = analytics.get("process_breakdown", [])
    source       = analytics.get("source", "Databricks")

    fta_current  = analytics.get("fta_current", 0.0)
    fta_prev     = analytics.get("fta_prev_week", fta_current)
    fta_rag      = analytics.get("fta_rag", "GREEN")
    fta_trend    = analytics.get("fta_trend", "stable")
    coq_current  = analytics.get("coq_current", 0.0)
    coq_rag      = analytics.get("coq_rag", "GREEN")
    eff_vs_bl    = analytics.get("efficiency_vs_bl")
    eff_rag      = analytics.get("efficiency_rag", "GREEN")

    # scope_mode
    pts = _s.get("databricks_process_types") or []
    scope_mode = _s.get("scope_mode") or ("defined" if pts else "all")

    # RAG colours (RGB)
    _C = {
        "GREEN":  RGBColor(0x15, 0x80, 0x3D),
        "AMBER":  RGBColor(0xB0, 0x60, 0x00),
        "RED":    RGBColor(0xC0, 0x39, 0x2B),
        "NAVY":   RGBColor(0x1A, 0x1A, 0x2E),
        "GRAY":   RGBColor(0x64, 0x74, 0x8B),
        "WHITE":  RGBColor(0xFF, 0xFF, 0xFF),
    }
    _BG = {
        "GREEN": "D5F5E3",
        "AMBER": "FEF3E2",
        "RED":   "FDE8E8",
    }

    overall = "GREEN"
    breaches = jira.get("sla_breaches", []) or []
    warnings = jira.get("sla_warnings", []) or []
    if breaches or fta_rag == "RED":
        overall = "RED"
    elif warnings or fta_rag == "AMBER" or coq_rag in ("AMBER", "RED"):
        overall = "AMBER"

    # ── Helpers ────────────────────────────────────────────────────────────────

    def _set_cell_bg(cell, hex_color: str):
        tc   = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd  = OxmlElement("w:shd")
        shd.set(qn("w:val"),   "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"),  hex_color)
        tcPr.append(shd)

    def _hdr_row(table, *cols, bg="1B2A4A"):
        row = table.rows[0]
        for i, text in enumerate(cols):
            cell = row.cells[i]
            _set_cell_bg(cell, bg)
            p    = cell.paragraphs[0]
            p.clear()
            run  = p.add_run(text)
            run.font.bold       = True
            run.font.size       = Pt(9)
            run.font.color.rgb  = _C["WHITE"]
            p.alignment         = WD_ALIGN_PARAGRAPH.CENTER

    def _data_row(table, row_idx, *vals, bg=None, align=WD_ALIGN_PARAGRAPH.CENTER):
        row = table.rows[row_idx]
        for i, text in enumerate(vals):
            cell = row.cells[i]
            if bg:
                _set_cell_bg(cell, bg)
            p    = cell.paragraphs[0]
            p.clear()
            run  = p.add_run(str(text))
            run.font.size = Pt(9)
            p.alignment   = align

    def _section_heading(doc, title: str, n: int):
        p   = doc.add_paragraph()
        run = p.add_run(f"{n}.  {title}")
        run.font.bold      = True
        run.font.size      = Pt(11)
        run.font.color.rgb = _C["NAVY"]
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after  = Pt(4)

    def _delta_str(curr: float, prev: float, higher_is_good: bool = True) -> str:
        d    = curr - prev
        sign = "+" if d >= 0 else ""
        dir_ = "up" if d > 0 else ("dn" if d < 0 else "->")
        return f"{sign}{d:.2f}%"

    def _eff_str(val) -> str:
        if val is None:
            return "N/A"
        diff = val - 100
        sign = "+" if diff >= 0 else ""
        return f"{val:.1f}% ({sign}{diff:.1f}%)"

    def _rag_prefix(rag: str) -> str:
        return {"RED": "[RED]", "AMBER": "[AMB]", "GREEN": "[OK] "}.get(rag, "     ")

    # ── Document setup ─────────────────────────────────────────────────────────

    doc = Document()
    for sect in doc.sections:
        sect.top_margin    = Cm(1.8)
        sect.bottom_margin = Cm(1.8)
        sect.left_margin   = Cm(2.0)
        sect.right_margin  = Cm(2.0)

    # ── Cover header ───────────────────────────────────────────────────────────

    title_p = doc.add_paragraph()
    title_r = title_p.add_run("Weekly Operational Report")
    title_r.font.bold      = True
    title_r.font.size      = Pt(18)
    title_r.font.color.rgb = _C["NAVY"]
    title_p.alignment      = WD_ALIGN_PARAGRAPH.LEFT

    sub_p = doc.add_paragraph()
    sub_r = sub_p.add_run(f"{proj_label}  |  Week of {week}  |  {date_str}")
    sub_r.font.size      = Pt(10)
    sub_r.font.color.rgb = _C["GRAY"]
    sub_p.paragraph_format.space_after = Pt(2)

    rag_p = doc.add_paragraph()
    rag_r = rag_p.add_run(f"Overall Status:  {overall}")
    rag_r.font.bold      = True
    rag_r.font.size      = Pt(11)
    rag_r.font.color.rgb = _C.get(overall, _C["GRAY"])
    rag_p.paragraph_format.space_after = Pt(6)

    # ── Auto-numbering ─────────────────────────────────────────────────────────

    sec_n = [0]

    def section(title):
        sec_n[0] += 1
        _section_heading(doc, title, sec_n[0])

    # ── Section 1: Executive Summary ───────────────────────────────────────────

    section("Executive Summary")

    headline = {
        "RED":   f"Immediate attention required — {project_name} quality metrics are below threshold.",
        "AMBER": f"{project_name} performing with caution — some metrics need monitoring this week.",
        "GREEN": f"{project_name} is on track — all key metrics within healthy range.",
    }[overall]
    hl_p = doc.add_paragraph()
    hl_r = hl_p.add_run(headline)
    hl_r.font.bold      = True
    hl_r.font.size      = Pt(10)
    hl_r.font.color.rgb = _C.get(overall, _C["GRAY"])
    hl_p.paragraph_format.space_after = Pt(4)

    def _bullet(text: str, rag: str = "GRAY", bold_prefix: str = ""):
        bp = doc.add_paragraph(style="List Bullet")
        if bold_prefix:
            br = bp.add_run(bold_prefix + ": ")
            br.font.bold = True
            br.font.size = Pt(9)
        vr = bp.add_run(text)
        vr.font.size      = Pt(9)
        vr.font.color.rgb = _C.get(rag, _C["GRAY"])
        bp.paragraph_format.space_after = Pt(2)

    trend_arrow = {"improving": "up", "stable": "->", "declining": "dn"}.get(fta_trend, "->")
    if "fta" in pm:
        _bullet(f"{fta_current:.2f}%  ({fta_rag})  trend {trend_arrow}", fta_rag, "FTA")
    if "efficiency" in pm:
        _bullet(f"{_eff_str(eff_vs_bl)}  ({eff_rag})", eff_rag, "Efficiency vs Baseline")
    if "coq" in pm or "copq" in pm:
        _bullet(f"{coq_current:.2f}%  ({coq_rag})", coq_rag, "Cost of Quality")
    if "yield" in pm:
        yv  = analytics.get("yield_current")
        yr  = analytics.get("yield_rag", "GREEN")
        yv_str = f"{yv:.2f}%" if yv is not None else "N/A"
        _bullet(f"{yv_str}  ({yr})", yr, "Yield (First Pass)")

    if breakdown:
        worst = min(breakdown, key=lambda p: p.get("fta_pct", 100))
        if worst.get("fta_rag") in ("AMBER", "RED"):
            _bullet(f'{worst["process_type"]}  at  {worst["fta_pct"]:.2f}%',
                    worst["fta_rag"], "Lowest FTA")

    if breaches:
        extra = f" + {len(warnings)} warning(s)" if warnings else ""
        _bullet(f"{len(breaches)} breach(es) requiring immediate follow-up{extra}", "RED", "Jira SLA")
    elif warnings:
        _bullet(f"{len(warnings)} warning(s) approaching breach threshold", "AMBER", "Jira SLA")
    else:
        _bullet("All clear — no breaches or warnings", "GREEN", "Jira SLA")

    _bullet(week, "GRAY", "Reporting Week")

    # ── Section 2: Key Performance Indicators ──────────────────────────────────

    section("Key Performance Indicators")

    kpi_cols = []
    if "fta" in pm:
        kpi_cols.append(("FTA — First Time Accuracy", fta_current, fta_prev, True, fta_rag))
    if "coq" in pm or "copq" in pm:
        prev_coq = m4w[1]["coq"] if len(m4w) > 1 else coq_current
        kpi_cols.append(("CoQ — Cost of Quality", coq_current, prev_coq, False, coq_rag))
    if "efficiency" in pm and eff_vs_bl is not None:
        kpi_cols.append(("Efficiency vs Baseline", eff_vs_bl, None, True, eff_rag))

    if kpi_cols:
        tbl = doc.add_table(rows=2, cols=len(kpi_cols))
        tbl.style = "Table Grid"
        _hdr_row(tbl, *[k[0] for k in kpi_cols])
        row = tbl.rows[1]
        for i, (label, curr, prev, hig, rag) in enumerate(kpi_cols):
            cell = row.cells[i]
            _set_cell_bg(cell, _BG.get(rag, "F1F5F9"))
            p    = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            val_r = p.add_run(f"{curr:.2f}%")
            val_r.font.bold      = True
            val_r.font.size      = Pt(14)
            val_r.font.color.rgb = _C.get(rag, _C["GRAY"])
            if prev is not None:
                d = curr - prev
                sign = "+" if d >= 0 else ""
                p.add_run(f"\n{sign}{d:.2f}% vs prior week").font.size = Pt(8)

    # ── Section 3: Process-type Breakdown ──────────────────────────────────────

    section(f"Process-type Breakdown — Week {week}")

    many = len(breakdown) > 15

    if scope_mode == "all" or (scope_mode.startswith("keyword") and many):
        display_rows = sorted(breakdown, key=lambda p: p.get("total_hrs", 0), reverse=True)[:10]
        mode_note = f"Top 10 process types by volume (of {len(breakdown)} total)"
    elif many:
        by_fta   = sorted(breakdown, key=lambda p: p.get("fta_pct", 100))
        bottom10 = by_fta[:10]
        top10    = list(reversed(by_fta[-10:]))
        display_rows = bottom10 + top10
        mode_note = f"Bottom 10 (worst FTA) + Top 10 (best FTA) of {len(breakdown)} total"
    else:
        display_rows = sorted(breakdown, key=lambda p: p.get("fta_pct", 100))
        mode_note = f"All {len(breakdown)} process type(s) — sorted by FTA (worst first)"

    note_p = doc.add_paragraph(mode_note)
    note_p.runs[0].font.size      = Pt(8)
    note_p.runs[0].font.color.rgb = _C["GRAY"]
    note_p.paragraph_format.space_after = Pt(3)

    if display_rows:
        tbl = doc.add_table(rows=1 + len(display_rows), cols=8)
        tbl.style = "Table Grid"
        _hdr_row(tbl, "Process Type", "FTA", "Eff vs BL", "CoQ", "Yield", "Sampling", "Tasks", "Hours")
        for ri, p in enumerate(display_rows):
            fr = p.get("fta_rag", "GREEN")
            row_bg = _BG.get(fr, "FFFFFF") if fr in ("RED", "AMBER") else None
            sp = p.get("sampling_pct")
            yp = p.get("yield_pct")
            _data_row(
                tbl, ri + 1,
                p.get("process_type", ""),
                f'{p.get("fta_pct", 0):.2f}%',
                (_eff_str(p.get("eff_vs_bl")) if p.get("has_prod", True) else "—"),
                f'{p.get("coq_pct", 0):.1f}%',
                f'{yp:.1f}%' if yp is not None else '—',
                f'{sp:.1f}%' if sp is not None else '—',
                f'{p.get("prod_tasks", 0):,}',
                f'{p.get("total_hrs", 0):.0f}h',
                bg=row_bg,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            # Color the FTA cell
            fta_cell = tbl.rows[ri + 1].cells[1]
            fta_cell.paragraphs[0].runs[0].font.color.rgb = _C.get(fr, _C["GRAY"])
            fta_cell.paragraphs[0].runs[0].font.bold = True

    # ── Section 3b: Operator Breakdown (only when the skill opts into operator level) ──

    ops = analytics.get("operator_breakdown") or []
    if ops:
        section(f"Operator Breakdown — Week {week}")
        note_p = doc.add_paragraph(
            f"{len(ops)} operator(s) with QC data, grouped by process type. "
            "FTA = first-pass acceptance rate per operator.")
        note_p.runs[0].font.size      = Pt(8)
        note_p.runs[0].font.color.rgb = _C["GRAY"]
        note_p.paragraph_format.space_after = Pt(3)

        tbl = doc.add_table(rows=1 + len(ops), cols=8)
        tbl.style = "Table Grid"
        _hdr_row(tbl, "Process Type", "Operator", "FTA", "Eff vs BL", "CoQ", "Sampling", "Tasks", "QC")
        for ri, o in enumerate(ops):
            fr = o.get("fta_rag", "GREEN")
            sp = o.get("sampling_pct")
            ev = o.get("eff_vs_bl")
            _data_row(
                tbl, ri + 1,
                o.get("process_type", ""),
                o.get("operator", ""),
                f'{o.get("fta_pct", 0):.2f}%',
                _eff_str(ev) if ev is not None else ("—" if not o.get("has_prod", True) else "N/A"),
                f'{o.get("coq_pct", 0):.1f}%',
                f'{sp:.1f}%' if sp is not None else '—',
                f'{o.get("prod_tasks", 0):,}',
                f'{o.get("total_qc", 0):,}',
                bg=_BG.get(fr, "FFFFFF") if fr in ("RED", "AMBER") else None,
            )
            fta_cell = tbl.rows[ri + 1].cells[2]
            fta_cell.paragraphs[0].runs[0].font.color.rgb = _C.get(fr, _C["GRAY"])
            fta_cell.paragraphs[0].runs[0].font.bold = True

    # ── Section 4: 4-Week Trend ────────────────────────────────────────────────

    if m4w:
        section(f"4-Week Trend — {project_name}")
        headers = ["Week"]
        if "fta" in pm:        headers.append("FTA")
        if "coq" in pm:        headers.append("CoQ")
        if "efficiency" in pm: headers.append("Efficiency vs BL")
        if "yield" in pm:      headers.append("Yield")
        headers.append("Sampling")

        tbl = doc.add_table(rows=1 + len(m4w), cols=len(headers))
        tbl.style = "Table Grid"
        _hdr_row(tbl, *headers)
        for ri, w in enumerate(m4w):
            is_latest = ri == len(m4w) - 1
            vals = [f"{w['week']}{' (latest)' if is_latest else ''}"]
            if "fta" in pm:
                fta_v = w.get("fta", 0.0)
                arr = ""
                if ri > 0:
                    d = fta_v - m4w[ri - 1].get("fta", fta_v)
                    arr = " up" if d > 0.5 else (" dn" if d < -0.5 else " ->")
                vals.append(f"{fta_v:.2f}%{arr}")
            if "coq" in pm:
                coq_v = w.get("coq", 0.0)
                arr = ""
                if ri > 0:
                    d = coq_v - m4w[ri - 1].get("coq", coq_v)
                    arr = " up" if d > 0.2 else (" dn" if d < -0.2 else " ->")
                vals.append(f"{coq_v:.2f}%{arr}")
            if "efficiency" in pm:
                vals.append(_eff_str(w.get("efficiency_vs_bl")))
            if "yield" in pm:
                yv = w.get("yield")
                vals.append(f"{yv:.2f}%" if yv is not None else "N/A")
            sp = w.get("sampling_pct")
            vals.append(f"{sp:.1f}%" if sp is not None else "—")
            bg = "EBF5FB" if is_latest else None
            _data_row(tbl, ri + 1, *vals, bg=bg)

        src_p = doc.add_paragraph(f"Source: {source}")
        src_p.runs[0].font.size      = Pt(8)
        src_p.runs[0].font.color.rgb = _C["GRAY"]

    # ── Section 5: Jira Sprint Health ─────────────────────────────────────────

    open_cnt = len(jira.get("open_tickets", []) or [])
    on_track = max(open_cnt - len(breaches) - len(warnings), 0)

    section(f"{jira_project} Sprint Health" if jira_project else "Sprint Health")
    sprint_tbl = doc.add_table(rows=2, cols=4)
    sprint_tbl.style = "Table Grid"
    _hdr_row(sprint_tbl, "Total Open", "SLA Breach", "SLA Warning", "On Track")
    for i, (val, rag_key) in enumerate([
        (str(open_cnt),       "GREEN"),
        (str(len(breaches)),  "RED"   if breaches else "GREEN"),
        (str(len(warnings)),  "AMBER" if warnings else "GREEN"),
        (str(on_track),       "GREEN"),
    ]):
        cell = sprint_tbl.rows[1].cells[i]
        _set_cell_bg(cell, _BG.get(rag_key, "F0FDF4"))
        p    = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r    = p.add_run(val)
        r.font.bold      = True
        r.font.size      = Pt(12)
        r.font.color.rgb = _C.get(rag_key, _C["GRAY"])

    if breaches:
        doc.add_paragraph()
        br_p = doc.add_paragraph("SLA Breaches")
        br_p.runs[0].font.bold      = True
        br_p.runs[0].font.color.rgb = _C["RED"]
        br_p.runs[0].font.size      = Pt(9)
        for t in breaches:
            bp = doc.add_paragraph(
                f"  {t.get('key', '')} — stale {t.get('hours_stale', '')}h"
                f"  ({t.get('assignee', '')})",
                style="List Bullet",
            )
            bp.runs[0].font.size = Pt(9)

    jira_src = jira.get("source", f"Jira {jira_project}")
    js_p = doc.add_paragraph(f"Source: {jira_src}")
    js_p.runs[0].font.size      = Pt(8)
    js_p.runs[0].font.color.rgb = _C["GRAY"]

    # ── Section 6: Confluence ─────────────────────────────────────────────────

    conf_pages = confluence.get("pages", []) or []
    if conf_pages:
        section("Confluence — Weekly Progress")
        for page in conf_pages[:5]:
            title   = page.get("title", "Untitled")
            updated = page.get("last_updated", "")
            author  = page.get("last_author", "")
            excerpt = page.get("excerpt", "")
            line = title
            if updated:
                line += f"  |  Updated {updated}"
            if author:
                line += f" by {author}"
            cp = doc.add_paragraph(style="List Bullet")
            cp.add_run(line).font.size = Pt(9)
            if excerpt:
                ep = doc.add_paragraph(f"    {excerpt[:150]}{'...' if len(excerpt) > 150 else ''}")
                ep.runs[0].font.size      = Pt(8)
                ep.runs[0].font.color.rgb = _C["GRAY"]

    # ── Section 7: Recommended Actions & Conclusion (P0 / P1 / P2) ────────────

    section("Recommended Actions & Conclusion")

    p0, p1, p2 = [], [], []

    for t in breaches:
        d = round(t.get("hours_stale", 0) / 24, 1)
        p0.append(
            f'Follow up with {t.get("assignee", "owner")} on {t.get("key", "")} '
            f'— silent {d} days, SLA breached.'
        )

    if "fta" in pm and fta_rag == "RED":
        p0.append(f"Overall FTA {fta_current:.2f}% is RED — investigate quality root causes immediately.")

    if "efficiency" in pm and eff_rag == "RED" and eff_vs_bl is not None:
        p0.append(
            f"Efficiency {_eff_str(eff_vs_bl)} is RED — identify throughput bottlenecks "
            f"in low-performing process types."
        )

    if "fta" in pm:
        for p in sorted(breakdown, key=lambda x: x.get("fta_pct", 100)):
            if p.get("fta_rag") == "RED":
                p1.append(
                    f'{p["process_type"]} — FTA {p["fta_pct"]:.2f}% (RED), '
                    f'CoQ {p.get("coq_pct", 0):.1f}%. Review rejection patterns.'
                )

    if "coq" in pm and coq_rag in ("RED_P1", "RED"):
        p1.append(f"CoQ {coq_current:.2f}% (RED) — audit rework root causes for week {week}.")

    if "fta" in pm and fta_rag == "AMBER":
        p2.append(
            f"Overall FTA {fta_current:.2f}% is AMBER — monitor trend; "
            f"cross-check quality-labelled tickets."
        )

    if "fta" in pm:
        for p in sorted(breakdown, key=lambda x: x.get("fta_pct", 100)):
            if p.get("fta_rag") == "AMBER":
                p2.append(
                    f'{p["process_type"]} — FTA {p["fta_pct"]:.2f}% (AMBER). '
                    f'Monitor next week; intervene if declining.'
                )

    if "coq" in pm and coq_rag in ("AMBER_P2", "AMBER", "AMBER_P3"):
        p2.append(f"CoQ {coq_current:.2f}% (AMBER) — investigate rework patterns.")

    if "efficiency" in pm and eff_rag == "AMBER" and eff_vs_bl is not None:
        p2.append(
            f"Efficiency {_eff_str(eff_vs_bl)} (AMBER) — review low-performing "
            f"process types before reaching RED."
        )

    for t in warnings:
        p2.append(
            f'Nudge {t.get("assignee", "owner")} on {t.get("key", "")} '
            f'— {t.get("hours_stale", "")}h without update, approaching SLA.'
        )

    if not p0 and not p1 and not p2:
        ok_p = doc.add_paragraph(style="List Bullet")
        ok_r = ok_p.add_run("[OK]  No immediate actions required — all metrics are on track.")
        ok_r.font.size      = Pt(9)
        ok_r.font.color.rgb = _C["GREEN"]
    else:
        idx = 1
        for tag, label, items, rag_key in [
            ("P0", "Immediate Action Required",            p0, "RED"),
            ("P1", "Fix This Week",                        p1, "RED"),
            ("P2", "Monitor — Intervene if Trend Continues", p2, "AMBER"),
        ]:
            if not items:
                continue
            grp_p = doc.add_paragraph()
            grp_r = grp_p.add_run(f"{tag} — {label}")
            grp_r.font.bold      = True
            grp_r.font.size      = Pt(10)
            grp_r.font.color.rgb = _C.get(rag_key, _C["GRAY"])
            grp_p.paragraph_format.space_before = Pt(8)
            grp_p.paragraph_format.space_after  = Pt(2)
            for item in items:
                ap = doc.add_paragraph(style="List Bullet")
                ar = ap.add_run(f"{idx}. {item}")
                ar.font.size      = Pt(9)
                ar.font.color.rgb = _C.get(rag_key, _C["GRAY"])
                ap.paragraph_format.space_after = Pt(3)
                idx += 1

        total = len(p0) + len(p1) + len(p2)
        parts = []
        if p0: parts.append(f"{len(p0)} immediate (P0)")
        if p1: parts.append(f"{len(p1)} high priority (P1)")
        if p2: parts.append(f"{len(p2)} to monitor (P2)")
        sum_p = doc.add_paragraph()
        sum_r = sum_p.add_run(
            f"Summary: {total} action{'s' if total != 1 else ''} — "
            + ", ".join(parts)
            + f". All actions require {user_name}'s approval before any external communication."
        )
        sum_r.font.size      = Pt(9)
        sum_r.font.color.rgb = _C["GRAY"]
        sum_p.paragraph_format.space_before = Pt(8)

    # ── Footer ─────────────────────────────────────────────────────────────────

    doc.add_paragraph()
    foot_p = doc.add_paragraph(
        f"UC3 Weekly Report  |  {jira_project}  |  Databricks  |  Confluence\n"
        f"Draft for {user_name}'s review before any distribution.\n"
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    foot_p.runs[0].font.size      = Pt(8)
    foot_p.runs[0].font.color.rgb = _C["GRAY"]

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
